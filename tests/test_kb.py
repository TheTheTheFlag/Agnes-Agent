"""回归测试：知识库管理（KB）——适配器层（文档/状态/分块映射、doc id 计算）与 API 端点。

全部走 mock：不初始化真实 LightRAG、不触网、不依赖 Neo4j/FAISS 进程。
"""
import asyncio
import importlib.util
import json
import os
import shutil
import tempfile
import unittest
from unittest import mock

from app.memory import ligraphrag_adapter as L


class _TmpWorkspace(unittest.TestCase):
    """把 _workspace_dir 指向临时目录，并写入 KV JSON 化的存储文件。"""

    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="kb_test_")
        self._pad = mock.patch.object(L, "_workspace_dir", return_value=self._tmp)
        self._pad.start()
        self.addCleanup(self._pad.stop)
        self.addCleanup(shutil.rmtree, self._tmp, ignore_errors=True)

    def _write(self, name, data):
        p = os.path.join(self._tmp, name)
        with open(p, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)


class DocIdTest(unittest.TestCase):
    def test_doc_id_md5_deterministic(self):
        a = L._doc_id_for_text("hello 知识库")
        self.assertTrue(a.startswith("doc-"))
        self.assertEqual(len(a), 4 + 32)
        self.assertEqual(a, L._doc_id_for_text("hello 知识库"))
        self.assertNotEqual(a, L._doc_id_for_text("hello 知识库."))


class DocumentsTest(_TmpWorkspace):
    def test_mapping_and_pagination(self):
        docs = {}
        for i in range(12):
            docs[f"doc-{i:032d}"] = {
                "status": "processed",
                "content_summary": f"文档{i}",
                "content_length": 100 + i,
                "file_path": "unknown_source",
                "chunks_count": i,
                "created_at": f"2026-09-14T00:00:0{i}+00:00",
                "updated_at": f"2026-09-14T00:{i:02d}:00+00:00",
            }
        self._write("kv_store_doc_status.json", docs)
        r = L.kb_documents("tid", page=1, page_size=10)
        self.assertEqual(r["total"], 12)
        self.assertEqual(len(r["docs"]), 10)
        # 按 updated_at 倒序：00:11 是 i=11 → 文档11 排最前（chunks_count=11）
        self.assertEqual(r["docs"][0]["title"], "unknown_source")
        self.assertEqual(r["docs"][0]["chunks_count"], 11)
        self.assertEqual(r["docs"][0]["status"], "processed")
        # 第二页只剩 2 条
        r2 = L.kb_documents("tid", page=2, page_size=10)
        self.assertEqual(len(r2["docs"]), 2)

    def test_missing_file_degrades(self):
        r = L.kb_documents("tid", page=1, page_size=50)
        self.assertEqual(r["total"], 0)
        self.assertEqual(r["docs"], [])

    def test_page_size_clamped(self):
        self._write("kv_store_doc_status.json", {
            "doc-" + "a" * 32: {"status": "processed", "updated_at": "x"},
        })
        self.assertEqual(L.kb_documents("tid", page_size=999999)["page_size"], 200)


class ChunksTest(_TmpWorkspace):
    def test_chunk_listing_sorted(self):
        self._write("kv_store_text_chunks.json", {
            "doc-1111-chunk-001": {"tokens": 30, "content": "B", "chunk_order_index": 1},
            "doc-1111-chunk-000": {"tokens": 20, "content": "A", "chunk_order_index": 0},
            "doc-2222-chunk-000": {"tokens": 5, "content": "other"},
        })
        r = L.kb_chunks("tid", "doc-1111")
        self.assertEqual([c["content"] for c in r["chunks"]], ["A", "B"])
        self.assertEqual(r["doc_id"], "doc-1111")

    def test_prefix_doc_id_normalized(self):
        r = L.kb_chunks("tid", "1111")
        self.assertEqual(r["doc_id"], "doc-1111")


class StatusTest(_TmpWorkspace):
    def test_status_aggregates(self):
        self._write("kv_store_doc_status.json", {
            "doc-a1": {"status": "processed"},
            "doc-a2": {"status": "processed"},
            "doc-a3": {"status": "failed"},
        })
        self._write("kv_store_text_chunks.json", {"k0": {}, "k1": {}})
        self._write("kv_store_full_entities.json", {"e0": {}})
        self._write("kv_store_full_relations.json", {"r0": {}})
        rag = mock.Mock()
        FaissCls = type("FaissVectorDBStorage", (), {})
        NeoCls = type("Neo4JStorage", (), {})
        rag.entities_vdb = FaissCls()
        rag.chunk_entity_relation_graph = NeoCls()
        with mock.patch.object(L, "get_lightrag", return_value=rag), \
                mock.patch.object(L, "_ensure_initialized"), \
                mock.patch.object(L, "graph_snapshot", return_value={"nodes": [1, 2], "edges": [1]}):
            st = L.kb_status("tid")
        self.assertEqual(st["doc_count"], 3)
        self.assertEqual(st["doc_statuses"]["processed"], 2)
        self.assertEqual(st["chunk_count"], 2)
        self.assertEqual(st["entity_count"], 1)
        self.assertEqual(st["relation_count"], 1)
        self.assertEqual(st["graph_nodes"], 2)
        self.assertEqual(st["vector_backend"], "faiss")
        self.assertEqual(st["graph_backend"], "neo4j")


class IngestDeleteTest(unittest.TestCase):
    def test_ingest_empty_rejected(self):
        r = L.kb_ingest("tid", [])
        self.assertFalse(r["ok"])
        self.assertIn("没有可喂入", r["error"])

    def test_ingest_forwards_to_insert(self):
        with mock.patch.object(L, "lightrag_insert") as ins:
            r = L.kb_ingest("tid", ["  ", "hello"])
        ins.assert_called_once_with("tid", ["hello"])
        self.assertTrue(r["ok"])
        self.assertEqual(r["doc_ids"], [L._doc_id_for_text("hello")])

    def test_ingest_failure_degrades(self):
        with mock.patch.object(L, "lightrag_insert", side_effect=RuntimeError("boom")):
            r = L.kb_ingest("tid", ["hello"])
        self.assertFalse(r["ok"])
        self.assertIn("boom", r["error"])


class KBAPITest(unittest.TestCase):
    """直接调用 handler 函数（绕过 HTTP/auth），mock 掉适配器函数。"""

    def setUp(self):
        self._p = mock.patch.object(L, "kb_threads", return_value=[])
        self._p.start()
        self.addCleanup(mock.patch.stopall)

    def test_threads_endpoint(self):
        from app.server.api.kb import kb_threads
        L.kb_threads.return_value = [{"id": "t1", "label": "t1", "dir": "/x"}]
        resp = asyncio.run(kb_threads())
        self.assertEqual(resp["threads"][0]["id"], "t1")

    def test_status_endpoint_forwards_thread(self):
        from app.server.api.kb import status
        with mock.patch.object(L, "kb_status", return_value={"doc_count": 2}):
            resp = asyncio.run(status(thread_id="abc"))
        self.assertEqual(resp["doc_count"], 2)

    def test_documents_endpoint_clamps_page(self):
        from app.server.api.kb import documents
        with mock.patch.object(L, "kb_documents") as doc:
            doc.return_value = {"docs": [], "total": 0}
            asyncio.run(documents(thread_id="t", page=3, page_size=25))
        self.assertEqual(doc.call_args[1]["page"], 3)

    def test_search_endpoint(self):
        from app.server.api.kb import search
        with mock.patch.object(L, "kb_search", return_value={"hybrid": "", "vector_hits": {}}) as s:
            resp = asyncio.run(search(thread_id="t", q="hello", top_k=6))
        self.assertEqual(resp["hybrid"], "")
        self.assertEqual(s.call_args[1]["top_k"], 6)

    def test_ingest_endpoint(self):
        from app.server.api.kb import ingest, IngestReq
        with mock.patch.object(L, "kb_ingest", return_value={"ok": True, "doc_ids": ["doc-x"]}) as k:
            resp = asyncio.run(ingest(IngestReq(thread_id="t", texts=["hello"])))
        self.assertTrue(resp["ok"])
        k.assert_called_once_with("t", ["hello"])

    def test_delete_endpoint(self):
        from app.server.api.kb import delete, DeleteReq
        with mock.patch.object(L, "kb_delete_doc", return_value={"ok": True, "doc_id": "doc-x"}) as k:
            resp = asyncio.run(delete(DeleteReq(thread_id="t", doc_id="doc-x")))
        self.assertTrue(resp["ok"])
        k.assert_called_once_with("t", "doc-x")

    def test_reprocess_endpoint(self):
        from app.server.api.kb import reprocess, ReprocessReq
        with mock.patch.object(L, "kb_reprocess", return_value={"ok": True}) as k:
            resp = asyncio.run(reprocess(ReprocessReq(thread_id="t", doc_id="doc-x")))
        self.assertTrue(resp["ok"])
        k.assert_called_once_with("t", "doc-x")

    def test_ingest_file_missing_file_404(self):
        from fastapi import HTTPException
        from app.server.api.kb import ingest_file, IngestFileReq
        with self.assertRaises(HTTPException) as cm:
            asyncio.run(ingest_file(IngestFileReq(thread_id="t", path="uploads/nope_123.md")))
        self.assertEqual(cm.exception.status_code, 404)

    def test_ingest_file_rejects_unsupported_ext(self):
        from fastapi import HTTPException
        from app.server.api.kb import ingest_file, IngestFileReq
        up_dir = os.path.dirname(os.path.abspath(__file__))
        fake = os.path.join(up_dir, "fake_kb_doc.exe")
        with open(fake, "w", encoding="utf-8") as f:
            f.write("x")
        self.addCleanup(os.unlink, fake)
        with mock.patch("app.server.api.kb.UPLOAD_DIR", os.path.dirname(fake)), \
                self.assertRaises(HTTPException) as cm:
            asyncio.run(ingest_file(IngestFileReq(thread_id="t", path="fake_kb_doc.exe")))
        self.assertEqual(cm.exception.status_code, 400)

    def test_ingest_file_pdf_extracts_text(self):
        import tempfile
        from app.server.api.kb import ingest_file, IngestFileReq

        # 手工构造一个极简合法 PDF（Helvetica 文本行 "Hello KB PDF"）
        def _minimal_pdf(text):
            content = ("BT /F1 24 Tf 72 700 Td (%s) Tj ET" % text.replace("(", "\\(").replace(")", "\\)")).encode("latin-1")
            objs = [
                b"<< /Type /Catalog /Pages 2 0 R >>",
                b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
                b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
                b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            ]
            out = bytearray(b"%PDF-1.4\n")
            offsets = []
            for i, body in enumerate(objs, 1):
                offsets.append(len(out))
                out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
            xref_pos = len(out)
            out += b"xref\n0 6\n0000000000 65535 f \n"
            for off in offsets:
                out += b"%010d 00000 n \n" % off
            out += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % xref_pos
            return bytes(out)

        with tempfile.TemporaryDirectory() as tmp:
            pdf = os.path.join(tmp, "sample.pdf")
            with open(pdf, "wb") as f:
                f.write(_minimal_pdf("Hello KB PDF"))
            with mock.patch("app.server.api.kb.UPLOAD_DIR", tmp), \
                    mock.patch.object(L, "kb_ingest", return_value={"ok": True, "doc_ids": ["d"]}) as k:
                resp = asyncio.run(ingest_file(IngestFileReq(thread_id="t", path="sample.pdf")))
            self.assertTrue(resp["ok"])
            self.assertTrue(k.called, "应把 PDF 文本喂入 kb_ingest")
            ingested = k.call_args.args[1][0]
            self.assertIn("Hello KB PDF", ingested)

    @unittest.skipUnless(
        importlib.util.find_spec("docx") and importlib.util.find_spec("markitdown"),
        "需要 python-docx 与 markitdown")
    def test_ingest_file_docx_converts_to_markdown(self):
        import tempfile
        from docx import Document
        from app.server.api.kb import ingest_file, IngestFileReq
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            doc.add_heading("月度总结", 0)
            doc.add_paragraph("核心指标继续保持增长，本月新增用户 1200 名。")
            fp = os.path.join(tmp, "report.docx")
            doc.save(fp)
            with mock.patch("app.server.api.kb.UPLOAD_DIR", tmp), \
                    mock.patch.object(L, "kb_ingest", return_value={"ok": True, "doc_ids": ["d"]}) as k:
                resp = asyncio.run(ingest_file(IngestFileReq(thread_id="t", path="report.docx")))
            self.assertTrue(resp["ok"])
            ingested = k.call_args.args[1][0]
            self.assertIn("月度总结", ingested)
            self.assertIn("1200", ingested)

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl") and importlib.util.find_spec("markitdown"),
        "需要 openpyxl 与 markitdown")
    def test_ingest_file_xlsx_converts_table(self):
        import tempfile
        from openpyxl import Workbook
        from app.server.api.kb import ingest_file, IngestFileReq
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            ws = wb.active
            ws.title = "销售"
            ws.append(["商品", "销量"])
            ws.append(["苹果", 10])
            ws.append(["香蕉", 6])
            fp = os.path.join(tmp, "sales.xlsx")
            wb.save(fp)
            with mock.patch("app.server.api.kb.UPLOAD_DIR", tmp), \
                    mock.patch.object(L, "kb_ingest", return_value={"ok": True, "doc_ids": ["d"]}) as k:
                resp = asyncio.run(ingest_file(IngestFileReq(thread_id="t", path="sales.xlsx")))
            self.assertTrue(resp["ok"])
            ingested = k.call_args.args[1][0]
            self.assertIn("苹果", ingested)
            self.assertIn("10", ingested)


if __name__ == "__main__":
    unittest.main()